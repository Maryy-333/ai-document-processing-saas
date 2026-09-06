import csv
import io
import uuid
from datetime import UTC, datetime, timedelta
from decimal import Decimal

import openpyxl
from fastapi.testclient import TestClient

from app.db.session import get_db
from app.main import app
from app.models import Document, Invoice, InvoiceLineItem, Organization, User
from app.models.enums import DocumentStatus
from tests.fixtures.auth import auth_headers


def make_org_user(db_session, org_name="Acme", email="user@example.com"):
    org = Organization(name=org_name)
    db_session.add(org)
    db_session.flush()
    user = User(organization_id=org.id, email=email, hashed_password="x")
    db_session.add(user)
    db_session.commit()
    return org, user


def make_document(db_session, org, user, status=DocumentStatus.APPROVED):
    doc = Document(
        organization_id=org.id,
        uploaded_by_user_id=user.id,
        original_filename="invoice.pdf",
        storage_key=f"{org.id}/{uuid.uuid4()}.pdf",
        content_type="application/pdf",
        file_size_bytes=100,
        status=status,
    )
    db_session.add(doc)
    db_session.commit()
    return doc


def make_approved_invoice(db_session, doc, user, **kwargs):
    defaults = dict(
        vendor_name="Acme Supplies",
        invoice_number="INV-1001",
        subtotal=Decimal("100.00"),
        tax=Decimal("10.00"),
        total=Decimal("110.00"),
        approved_at=datetime.now(UTC),
        approved_by_user_id=user.id,
    )
    defaults.update(kwargs)
    invoice = Invoice(document_id=doc.id, **defaults)
    db_session.add(invoice)
    db_session.commit()
    db_session.refresh(invoice)
    return invoice


def add_line_item(db_session, invoice, order=0, **kwargs):
    defaults = dict(
        description="Widget",
        quantity=Decimal("1"),
        unit_price=Decimal("10.00"),
        amount=Decimal("10.00"),
    )
    defaults.update(kwargs)
    item = InvoiceLineItem(invoice_id=invoice.id, line_order=order, **defaults)
    db_session.add(item)
    db_session.commit()
    return item


def client_with_db(db_session) -> TestClient:
    def _get_db_override():
        yield db_session

    app.dependency_overrides[get_db] = _get_db_override
    return TestClient(app)


def teardown_overrides():
    app.dependency_overrides.clear()


# ---------- Single invoice CSV ----------


def test_single_invoice_csv_export_success(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user)
    invoice = make_approved_invoice(db_session, doc, user)
    add_line_item(db_session, invoice)

    client = client_with_db(db_session)
    try:
        response = client.get(
            f"/api/v1/documents/{doc.id}/export/csv", headers=auth_headers(user.id)
        )
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("text/csv")
        assert f"invoice-{invoice.id}.csv" in response.headers["content-disposition"]

        reader = csv.DictReader(io.StringIO(response.text))
        rows = list(reader)
        assert len(rows) == 1
        assert rows[0]["vendor_name"] == "Acme Supplies"
        assert rows[0]["invoice_number"] == "INV-1001"
        assert rows[0]["subtotal"] == "100.00"
        assert rows[0]["line_item_description"] == "Widget"
    finally:
        teardown_overrides()


def test_single_invoice_csv_zero_line_items_still_exports(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user)
    make_approved_invoice(db_session, doc, user)

    client = client_with_db(db_session)
    try:
        response = client.get(
            f"/api/v1/documents/{doc.id}/export/csv", headers=auth_headers(user.id)
        )
        assert response.status_code == 200
        reader = csv.DictReader(io.StringIO(response.text))
        rows = list(reader)
        assert len(rows) == 1
        assert rows[0]["line_item_description"] == ""
    finally:
        teardown_overrides()


def test_single_invoice_csv_multiple_line_items(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user)
    invoice = make_approved_invoice(db_session, doc, user)
    add_line_item(db_session, invoice, 0, description="A")
    add_line_item(db_session, invoice, 1, description="B")

    client = client_with_db(db_session)
    try:
        response = client.get(
            f"/api/v1/documents/{doc.id}/export/csv", headers=auth_headers(user.id)
        )
        reader = csv.DictReader(io.StringIO(response.text))
        rows = list(reader)
        assert len(rows) == 2
        assert [r["line_item_description"] for r in rows] == ["A", "B"]
    finally:
        teardown_overrides()


# ---------- Single invoice XLSX ----------


def test_single_invoice_xlsx_export_success(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user)
    invoice = make_approved_invoice(db_session, doc, user)
    add_line_item(db_session, invoice)

    client = client_with_db(db_session)
    try:
        response = client.get(
            f"/api/v1/documents/{doc.id}/export/xlsx", headers=auth_headers(user.id)
        )
        assert response.status_code == 200
        assert response.headers["content-type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert f"invoice-{invoice.id}.xlsx" in response.headers["content-disposition"]

        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        assert rows[0][0] == "invoice_id"
        assert rows[1][2] == "Acme Supplies"
    finally:
        teardown_overrides()


# ---------- State protection ----------


def test_export_review_required_blocked(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user, status=DocumentStatus.REVIEW_REQUIRED)
    db_session.add(Invoice(document_id=doc.id, vendor_name="Draft"))
    db_session.commit()

    client = client_with_db(db_session)
    try:
        response = client.get(
            f"/api/v1/documents/{doc.id}/export/csv", headers=auth_headers(user.id)
        )
        assert response.status_code == 409
    finally:
        teardown_overrides()


def test_export_rejected_blocked(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user, status=DocumentStatus.REJECTED)
    db_session.add(Invoice(document_id=doc.id, vendor_name="Rejected", rejection_reason="bad"))
    db_session.commit()

    client = client_with_db(db_session)
    try:
        response = client.get(
            f"/api/v1/documents/{doc.id}/export/csv", headers=auth_headers(user.id)
        )
        assert response.status_code == 409
    finally:
        teardown_overrides()


def test_export_uploaded_blocked(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user, status=DocumentStatus.UPLOADED)

    client = client_with_db(db_session)
    try:
        response = client.get(
            f"/api/v1/documents/{doc.id}/export/csv", headers=auth_headers(user.id)
        )
        assert response.status_code == 409
    finally:
        teardown_overrides()


def test_export_text_extracted_blocked(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user, status=DocumentStatus.TEXT_EXTRACTED)

    client = client_with_db(db_session)
    try:
        response = client.get(
            f"/api/v1/documents/{doc.id}/export/csv", headers=auth_headers(user.id)
        )
        assert response.status_code == 409
    finally:
        teardown_overrides()


def test_export_failed_blocked(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user, status=DocumentStatus.FAILED)

    client = client_with_db(db_session)
    try:
        response = client.get(
            f"/api/v1/documents/{doc.id}/export/csv", headers=auth_headers(user.id)
        )
        assert response.status_code == 409
    finally:
        teardown_overrides()


def test_repeated_export_does_not_mutate_state(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user)
    invoice = make_approved_invoice(db_session, doc, user)

    client = client_with_db(db_session)
    try:
        client.get(f"/api/v1/documents/{doc.id}/export/csv", headers=auth_headers(user.id))
        client.get(f"/api/v1/documents/{doc.id}/export/xlsx", headers=auth_headers(user.id))
        client.get(f"/api/v1/documents/{doc.id}/export/csv", headers=auth_headers(user.id))

        db_session.refresh(doc)
        db_session.refresh(invoice)
        assert doc.status == DocumentStatus.APPROVED
        assert invoice.approved_by_user_id == user.id
    finally:
        teardown_overrides()


# ---------- Organization isolation ----------


def test_cannot_export_another_organizations_invoice(db_session):
    org_a, user_a = make_org_user(db_session, "Org A", "a@example.com")
    _, user_b = make_org_user(db_session, "Org B", "b@example.com")
    doc = make_document(db_session, org_a, user_a)
    make_approved_invoice(db_session, doc, user_a)

    client = client_with_db(db_session)
    try:
        response = client.get(
            f"/api/v1/documents/{doc.id}/export/csv", headers=auth_headers(user_b.id)
        )
        assert response.status_code == 404
    finally:
        teardown_overrides()


def test_organization_export_excludes_other_orgs_invoices(db_session):
    org_a, user_a = make_org_user(db_session, "Org A", "a@example.com")
    org_b, user_b = make_org_user(db_session, "Org B", "b@example.com")
    doc_a = make_document(db_session, org_a, user_a)
    make_approved_invoice(db_session, doc_a, user_a, invoice_number="ORG-A-INV")
    doc_b = make_document(db_session, org_b, user_b)
    make_approved_invoice(db_session, doc_b, user_b, invoice_number="ORG-B-INV")

    client = client_with_db(db_session)
    try:
        response = client.get("/api/v1/documents/export/csv", headers=auth_headers(user_a.id))
        text = response.text
        assert "ORG-A-INV" in text
        assert "ORG-B-INV" not in text
    finally:
        teardown_overrides()


# ---------- Authentication ----------


def test_export_missing_auth_returns_401(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user)
    make_approved_invoice(db_session, doc, user)

    client = client_with_db(db_session)
    try:
        response = client.get(f"/api/v1/documents/{doc.id}/export/csv")
        assert response.status_code == 401
    finally:
        teardown_overrides()


def test_export_invalid_token_returns_401(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user)
    make_approved_invoice(db_session, doc, user)

    client = client_with_db(db_session)
    try:
        response = client.get(
            f"/api/v1/documents/{doc.id}/export/csv",
            headers={"Authorization": "Bearer invalid.token.here"},
        )
        assert response.status_code == 401
    finally:
        teardown_overrides()


def test_organization_export_missing_auth_returns_401(db_session):
    client = client_with_db(db_session)
    try:
        response = client.get("/api/v1/documents/export/csv")
        assert response.status_code == 401
    finally:
        teardown_overrides()


def test_export_inactive_user_returns_401(db_session):
    from app.core.config import get_settings
    from app.core.security import create_access_token

    org = Organization(name="Inactive Org")
    db_session.add(org)
    db_session.flush()
    user = User(
        organization_id=org.id, email="inactive@example.com", hashed_password="x", is_active=False
    )
    db_session.add(user)
    db_session.commit()
    doc = make_document(db_session, org, user)
    make_approved_invoice(db_session, doc, user)

    settings = get_settings()
    token = create_access_token(
        user.id,
        secret_key=settings.jwt_secret_key,
        algorithm=settings.jwt_algorithm,
        expire_minutes=30,
    )

    client = client_with_db(db_session)
    try:
        response = client.get(
            f"/api/v1/documents/{doc.id}/export/csv",
            headers={"Authorization": f"Bearer {token}"},
        )
        assert response.status_code == 401
    finally:
        teardown_overrides()


# ---------- Organization-wide export ----------


def test_organization_csv_export_multiple_invoices(db_session):
    org, user = make_org_user(db_session)
    for i in range(3):
        doc = make_document(db_session, org, user)
        make_approved_invoice(db_session, doc, user, invoice_number=f"INV-{i}")

    client = client_with_db(db_session)
    try:
        response = client.get("/api/v1/documents/export/csv", headers=auth_headers(user.id))
        assert response.status_code == 200
        reader = csv.DictReader(io.StringIO(response.text))
        rows = list(reader)
        assert len(rows) == 3
    finally:
        teardown_overrides()


def test_organization_csv_export_excludes_non_approved(db_session):
    org, user = make_org_user(db_session)
    approved_doc = make_document(db_session, org, user, status=DocumentStatus.APPROVED)
    make_approved_invoice(db_session, approved_doc, user, invoice_number="APPROVED-ONE")
    review_doc = make_document(db_session, org, user, status=DocumentStatus.REVIEW_REQUIRED)
    db_session.add(
        Invoice(document_id=review_doc.id, vendor_name="Draft", invoice_number="DRAFT-ONE")
    )
    db_session.commit()

    client = client_with_db(db_session)
    try:
        response = client.get("/api/v1/documents/export/csv", headers=auth_headers(user.id))
        assert "APPROVED-ONE" in response.text
        assert "DRAFT-ONE" not in response.text
    finally:
        teardown_overrides()


def test_organization_export_zero_invoices_returns_valid_empty_csv(db_session):
    org, user = make_org_user(db_session)
    client = client_with_db(db_session)
    try:
        response = client.get("/api/v1/documents/export/csv", headers=auth_headers(user.id))
        assert response.status_code == 200
        reader = csv.reader(io.StringIO(response.text))
        header = next(reader)
        assert header[0] == "invoice_id"
    finally:
        teardown_overrides()


def test_organization_export_deterministic_ordering(db_session):
    org, user = make_org_user(db_session)
    base = datetime.now(UTC)
    doc1 = make_document(db_session, org, user)
    make_approved_invoice(
        db_session, doc1, user, invoice_number="SECOND", approved_at=base + timedelta(seconds=5)
    )
    doc2 = make_document(db_session, org, user)
    make_approved_invoice(db_session, doc2, user, invoice_number="FIRST", approved_at=base)

    client = client_with_db(db_session)
    try:
        response = client.get("/api/v1/documents/export/csv", headers=auth_headers(user.id))
        reader = csv.DictReader(io.StringIO(response.text))
        rows = list(reader)
        assert [r["invoice_number"] for r in rows] == ["FIRST", "SECOND"]
    finally:
        teardown_overrides()


def test_organization_csv_and_xlsx_equivalent_data(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user)
    make_approved_invoice(db_session, doc, user, invoice_number="EQUIV-TEST")

    client = client_with_db(db_session)
    try:
        csv_response = client.get("/api/v1/documents/export/csv", headers=auth_headers(user.id))
        xlsx_response = client.get(
            "/api/v1/documents/export/xlsx", headers=auth_headers(user.id)
        )

        csv_rows = list(csv.DictReader(io.StringIO(csv_response.text)))
        workbook = openpyxl.load_workbook(io.BytesIO(xlsx_response.content))
        sheet = workbook.active
        xlsx_data_rows = list(sheet.iter_rows(min_row=2, values_only=True))

        assert len(csv_rows) == len(xlsx_data_rows) == 1
        assert csv_rows[0]["invoice_number"] == "EQUIV-TEST"
        assert xlsx_data_rows[0][5] == "EQUIV-TEST"
    finally:
        teardown_overrides()
