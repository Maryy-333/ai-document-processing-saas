import csv
import io
import uuid
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal

import openpyxl
import pytest

from app.core.exceptions import NotFoundError
from app.models import Document, Invoice, InvoiceLineItem, Organization, User
from app.models.enums import DocumentStatus
from app.services.export_service import (
    InvoiceNotApprovedError,
    export_organization_invoices,
    export_single_invoice,
    flatten_invoices_to_rows,
    generate_csv,
    generate_xlsx,
    get_approved_invoice_for_export,
    list_organization_approved_invoices,
)


def make_org_user(db_session, org_name="Acme", email="user@example.com"):
    org = Organization(name=org_name)
    db_session.add(org)
    db_session.flush()
    user = User(organization_id=org.id, email=email, hashed_password="x")
    db_session.add(user)
    db_session.commit()
    return org, user


def make_document(db_session, org, user, status=DocumentStatus.APPROVED):
    doc = Document(
        organization_id=org.id,
        uploaded_by_user_id=user.id,
        original_filename="invoice.pdf",
        storage_key=f"{org.id}/{uuid.uuid4()}.pdf",
        content_type="application/pdf",
        file_size_bytes=100,
        status=status,
    )
    db_session.add(doc)
    db_session.commit()
    return doc


def make_approved_invoice(db_session, doc, user, **kwargs):
    defaults = dict(
        vendor_name="Acme Supplies",
        vendor_address="123 Main St",
        customer_name="Example Co",
        invoice_number="INV-1001",
        currency="USD",
        subtotal=Decimal("100.00"),
        tax=Decimal("10.00"),
        total=Decimal("110.00"),
        payment_terms="Net 30",
        approved_at=datetime.now(UTC),
        approved_by_user_id=user.id,
    )
    defaults.update(kwargs)
    invoice = Invoice(document_id=doc.id, **defaults)
    db_session.add(invoice)
    db_session.commit()
    db_session.refresh(invoice)
    return invoice


def add_line_item(db_session, invoice, order, **kwargs):
    defaults = dict(
        description="Widget",
        quantity=Decimal("1"),
        unit_price=Decimal("10.00"),
        amount=Decimal("10.00"),
    )
    defaults.update(kwargs)
    item = InvoiceLineItem(invoice_id=invoice.id, line_order=order, **defaults)
    db_session.add(item)
    db_session.commit()
    return item


# ---------- Retrieval / state gating ----------


def test_get_approved_invoice_for_export_success(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user, status=DocumentStatus.APPROVED)
    invoice = make_approved_invoice(db_session, doc, user)

    result = get_approved_invoice_for_export(db_session, doc.id, org.id)
    assert result.id == invoice.id


@pytest.mark.parametrize(
    "status",
    [
        DocumentStatus.UPLOADED,
        DocumentStatus.TEXT_EXTRACTED,
        DocumentStatus.REVIEW_REQUIRED,
        DocumentStatus.REJECTED,
        DocumentStatus.FAILED,
    ],
)
def test_export_blocked_for_non_approved_states(db_session, status):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user, status=status)
    with pytest.raises(InvoiceNotApprovedError):
        get_approved_invoice_for_export(db_session, doc.id, org.id)


def test_export_cross_org_raises_not_found(db_session):
    org_a, user_a = make_org_user(db_session, "Org A", "a@example.com")
    org_b, _ = make_org_user(db_session, "Org B", "b@example.com")
    doc = make_document(db_session, org_a, user_a, status=DocumentStatus.APPROVED)
    make_approved_invoice(db_session, doc, user_a)

    with pytest.raises(NotFoundError):
        get_approved_invoice_for_export(db_session, doc.id, org_b.id)


def test_export_does_not_mutate_document_or_invoice_state(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user, status=DocumentStatus.APPROVED)
    invoice = make_approved_invoice(db_session, doc, user)

    export_single_invoice(db_session, doc.id, org.id, fmt="csv")
    export_single_invoice(db_session, doc.id, org.id, fmt="xlsx")

    db_session.refresh(doc)
    db_session.refresh(invoice)
    assert doc.status == DocumentStatus.APPROVED
    assert invoice.approved_at is not None


# ---------- Flattening ----------


def test_flatten_invoice_with_no_line_items_produces_one_row(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user)
    invoice = make_approved_invoice(db_session, doc, user)

    rows = flatten_invoices_to_rows([invoice])
    assert len(rows) == 1
    assert rows[0]["line_item_description"] is None
    assert rows[0]["line_item_amount"] is None
    assert rows[0]["vendor_name"] == "Acme Supplies"


def test_flatten_invoice_with_multiple_line_items_produces_multiple_rows(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user)
    invoice = make_approved_invoice(db_session, doc, user)
    add_line_item(db_session, invoice, 0, description="First")
    add_line_item(db_session, invoice, 1, description="Second")
    add_line_item(db_session, invoice, 2, description="Third")
    db_session.refresh(invoice)

    rows = flatten_invoices_to_rows([invoice])
    assert len(rows) == 3
    assert [r["line_item_description"] for r in rows] == ["First", "Second", "Third"]
    assert all(r["vendor_name"] == "Acme Supplies" for r in rows)


def test_flatten_preserves_decimal_precision(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user)
    invoice = make_approved_invoice(
        db_session, doc, user, subtotal=Decimal("1234.56"), tax=Decimal("98.76")
    )
    rows = flatten_invoices_to_rows([invoice])
    assert rows[0]["subtotal"] == Decimal("1234.56")
    assert rows[0]["tax"] == Decimal("98.76")


# ---------- CSV generation ----------


def test_generate_csv_has_correct_headers():
    csv_bytes = generate_csv([])
    text = csv_bytes.decode("utf-8")
    reader = csv.reader(io.StringIO(text))
    header = next(reader)
    assert header == [
        "invoice_id",
        "document_id",
        "vendor_name",
        "vendor_address",
        "customer_name",
        "invoice_number",
        "invoice_date",
        "due_date",
        "currency",
        "subtotal",
        "tax",
        "total",
        "payment_terms",
        "approved_at",
        "approved_by_user_id",
        "line_item_description",
        "line_item_quantity",
        "line_item_unit_price",
        "line_item_amount",
    ]


def test_generate_csv_zero_rows_still_has_header():
    csv_bytes = generate_csv([])
    text = csv_bytes.decode("utf-8")
    lines = text.strip("\r\n").split("\r\n")
    assert len(lines) == 1


def test_generate_csv_decimal_not_converted_to_float(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user)
    invoice = make_approved_invoice(db_session, doc, user, subtotal=Decimal("100.10"))
    rows = flatten_invoices_to_rows([invoice])
    csv_bytes = generate_csv(rows)
    text = csv_bytes.decode("utf-8")
    assert "100.10" in text
    assert "100.099999" not in text


def test_generate_csv_handles_commas_quotes_newlines(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user)
    tricky_name = 'Vendor, "Special" Inc.\nSecond line'
    invoice = make_approved_invoice(db_session, doc, user, vendor_name=tricky_name)
    rows = flatten_invoices_to_rows([invoice])
    csv_bytes = generate_csv(rows)
    text = csv_bytes.decode("utf-8")

    reader = csv.reader(io.StringIO(text))
    next(reader)
    data_row = next(reader)
    assert data_row[2] == tricky_name


def test_generate_csv_none_values_are_blank(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user)
    invoice = make_approved_invoice(db_session, doc, user, due_date=None, payment_terms=None)
    rows = flatten_invoices_to_rows([invoice])
    csv_bytes = generate_csv(rows)
    text = csv_bytes.decode("utf-8")
    reader = csv.DictReader(io.StringIO(text))
    row = next(reader)
    assert row["due_date"] == ""
    assert row["payment_terms"] == ""


def test_generate_csv_date_representation_is_deterministic(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user)
    invoice = make_approved_invoice(db_session, doc, user, invoice_date=date(2026, 8, 20))
    rows = flatten_invoices_to_rows([invoice])
    csv_bytes = generate_csv(rows)
    text = csv_bytes.decode("utf-8")
    assert "2026-08-20" in text


# ---------- XLSX generation ----------


def test_generate_xlsx_is_valid_workbook():
    xlsx_bytes = generate_xlsx([])
    workbook = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert workbook.active is not None


def test_generate_xlsx_has_correct_headers():
    xlsx_bytes = generate_xlsx([])
    workbook = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    sheet = workbook.active
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert header[0] == "invoice_id"
    assert header[-1] == "line_item_amount"
    assert len(header) == 19


def test_generate_xlsx_matches_csv_data(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user)
    invoice = make_approved_invoice(db_session, doc, user)
    add_line_item(db_session, invoice, 0)
    db_session.refresh(invoice)
    rows = flatten_invoices_to_rows([invoice])

    csv_bytes = generate_csv(rows)
    xlsx_bytes = generate_xlsx(rows)

    csv_text = csv_bytes.decode("utf-8")
    csv_reader = csv.reader(io.StringIO(csv_text))
    csv_header = next(csv_reader)
    csv_row = next(csv_reader)

    workbook = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    sheet = workbook.active
    xlsx_rows = list(sheet.iter_rows(values_only=True))
    xlsx_header = list(xlsx_rows[0])
    xlsx_row = [str(v) if v is not None else "" for v in xlsx_rows[1]]

    assert csv_header == xlsx_header
    assert csv_row == xlsx_row


def test_generate_xlsx_decimal_not_stored_as_float(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user)
    invoice = make_approved_invoice(db_session, doc, user, total=Decimal("100.10"))
    rows = flatten_invoices_to_rows([invoice])
    xlsx_bytes = generate_xlsx(rows)
    workbook = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    sheet = workbook.active
    total_col_index = 11  # 0-based: total is the 12th column — see _COLUMNS order
    value = list(sheet.iter_rows(values_only=True))[1][total_col_index]
    assert value == "100.10"
    assert not isinstance(value, float)


# ---------- Organization-wide export ----------


def test_list_organization_approved_invoices_excludes_non_approved(db_session):
    org, user = make_org_user(db_session)
    approved_doc = make_document(db_session, org, user, status=DocumentStatus.APPROVED)
    make_approved_invoice(db_session, approved_doc, user)

    review_doc = make_document(db_session, org, user, status=DocumentStatus.REVIEW_REQUIRED)
    db_session.add(Invoice(document_id=review_doc.id, vendor_name="Not approved"))
    db_session.commit()

    results = list_organization_approved_invoices(db_session, org.id)
    assert len(results) == 1
    assert results[0].document_id == approved_doc.id


def test_list_organization_approved_invoices_excludes_other_orgs(db_session):
    org_a, user_a = make_org_user(db_session, "Org A", "a@example.com")
    org_b, user_b = make_org_user(db_session, "Org B", "b@example.com")
    doc_a = make_document(db_session, org_a, user_a, status=DocumentStatus.APPROVED)
    make_approved_invoice(db_session, doc_a, user_a)
    doc_b = make_document(db_session, org_b, user_b, status=DocumentStatus.APPROVED)
    make_approved_invoice(db_session, doc_b, user_b)

    results_a = list_organization_approved_invoices(db_session, org_a.id)
    assert len(results_a) == 1
    assert results_a[0].document_id == doc_a.id


def test_list_organization_approved_invoices_deterministic_ordering(db_session):
    org, user = make_org_user(db_session)
    base = datetime.now(UTC)

    doc1 = make_document(db_session, org, user, status=DocumentStatus.APPROVED)
    inv1 = make_approved_invoice(db_session, doc1, user, approved_at=base + timedelta(seconds=2))
    doc2 = make_document(db_session, org, user, status=DocumentStatus.APPROVED)
    inv2 = make_approved_invoice(db_session, doc2, user, approved_at=base)
    doc3 = make_document(db_session, org, user, status=DocumentStatus.APPROVED)
    inv3 = make_approved_invoice(db_session, doc3, user, approved_at=base + timedelta(seconds=1))

    results = list_organization_approved_invoices(db_session, org.id)
    assert [r.id for r in results] == [inv2.id, inv3.id, inv1.id]


def test_organization_export_zero_approved_invoices_returns_valid_empty_export(db_session):
    org, user = make_org_user(db_session)
    export_file = export_organization_invoices(db_session, org.id, fmt="csv")
    text = export_file.content.decode("utf-8")
    reader = csv.reader(io.StringIO(text))
    header = next(reader)
    assert header[0] == "invoice_id"
    with pytest.raises(StopIteration):
        next(reader)


def test_organization_export_xlsx_zero_approved_invoices_valid_workbook(db_session):
    org, user = make_org_user(db_session)
    export_file = export_organization_invoices(db_session, org.id, fmt="xlsx")
    workbook = openpyxl.load_workbook(io.BytesIO(export_file.content))
    assert workbook.active.max_row == 1


def test_organization_export_includes_multiple_approved_invoices(db_session):
    org, user = make_org_user(db_session)
    for i in range(3):
        doc = make_document(db_session, org, user, status=DocumentStatus.APPROVED)
        make_approved_invoice(db_session, doc, user, invoice_number=f"INV-{i}")

    export_file = export_organization_invoices(db_session, org.id, fmt="csv")
    text = export_file.content.decode("utf-8")
    reader = csv.DictReader(io.StringIO(text))
    rows = list(reader)
    assert len(rows) == 3
    assert {r["invoice_number"] for r in rows} == {"INV-0", "INV-1", "INV-2"}


# ---------- Filenames / media types ----------


def test_single_export_filename_and_media_type_csv(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user)
    invoice = make_approved_invoice(db_session, doc, user)

    export_file = export_single_invoice(db_session, doc.id, org.id, fmt="csv")
    assert export_file.filename == f"invoice-{invoice.id}.csv"
    assert export_file.media_type == "text/csv"


def test_single_export_filename_and_media_type_xlsx(db_session):
    org, user = make_org_user(db_session)
    doc = make_document(db_session, org, user)
    invoice = make_approved_invoice(db_session, doc, user)

    export_file = export_single_invoice(db_session, doc.id, org.id, fmt="xlsx")
    assert export_file.filename == f"invoice-{invoice.id}.xlsx"
    assert (
        export_file.media_type
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
